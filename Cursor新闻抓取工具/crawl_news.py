#!/usr/bin/env python3
"""按 Excel 行序提取公开新闻；支持站点配置或通用模式，不覆盖已有 TXT。"""
from __future__ import annotations

import argparse
import csv
import hashlib
import json
import re
import sys
import time
import urllib.error
import urllib.parse
import urllib.request
from collections import Counter
from datetime import datetime
from pathlib import Path

from openpyxl import load_workbook
from extractor import ExtractionError, extract_article, extract_generic_article
from public_http import public_opener

HERE = Path(__file__).resolve().parent
ID_RE = re.compile(r"^[0-9a-f]{16}$")
TRACKING = {"fbclid", "gclid", "mc_cid", "mc_eid"}


def normalize_url(url):
    try:
        p = urllib.parse.urlsplit(str(url).strip())
        if p.scheme.lower() not in {"https", "http"} or not p.hostname or p.username or p.password:
            return ""
        if p.port not in (None, 80, 443):
            return ""
        query = [(k, v) for k, v in urllib.parse.parse_qsl(p.query, keep_blank_values=True)
                 if not k.lower().startswith("utm_") and k.lower() not in TRACKING]
        return urllib.parse.urlunsplit((p.scheme.lower(), p.netloc.lower(), p.path or "/",
                                      urllib.parse.urlencode(sorted(query)), ""))
    except ValueError:
        return ""


def read_rows(path, sheet="headline_urls"):
    book = load_workbook(path, read_only=True, data_only=True)
    try:
        if sheet not in book.sheetnames:
            raise ValueError("工作表不存在：" + sheet)
        iterator = book[sheet].iter_rows(values_only=True)
        headers = [str(v or "").strip() for v in next(iterator)]
        if not {"headline_id", "headline", "url"}.issubset(headers):
            raise ValueError("工作表必须包含 headline_id、headline、url 列")
        result = []
        for row_number, values in enumerate(iterator, 2):
            if not any(v is not None for v in values):
                continue
            record = dict(zip(headers, values))
            record["excel_row"] = row_number
            for k in ("headline_id", "headline", "url"):
                record[k] = str(record.get(k) or "").strip()
            result.append(record)
        return result
    finally:
        book.close()


def body_from_saved(text):
    # 兼容已有三批 TXT：元数据在文首或文末；不能直接在第一个 Source: 截断。
    lines = text.split("\n---")[0].splitlines()
    clean = []
    for i, line in enumerate(lines):
        s = line.strip()
        if i == 0 or re.match(r"^(Headline ID|Source(?: URL)?|Author|Publisher|Published(?: \(.*?\))?|Updated|License|Terms of use|Rights|Downloaded|Retrieved):", s):
            continue
        if "©" in s or s.startswith(("Copyright ", "Personal, non-commercial study copy.", "Single local copy for personal")):
            continue
        if s.startswith("Accuracy and availability of transcripts"):
            break
        if s:
            clean.append(s)
    return "\n".join(clean)


def existing_state(project, output, rows):
    paths = set()
    for directory in (project / "Current Output" / "txt", project / "New Output", output):
        if directory.exists():
            paths.update(directory.glob("*.txt"))
    ids, urls, bodies = set(), set(), {}
    for path in sorted(paths):
        if not ID_RE.fullmatch(path.stem):
            continue
        text = path.read_text(encoding="utf-8-sig")
        ids.add(path.stem)
        bodies[path.stem] = body_from_saved(text)
        for match in re.finditer(r"^Source(?: URL)?:\s*(https?://\S+)", text, re.M):
            if normalize_url(match.group(1)):
                urls.add(normalize_url(match.group(1)))
    # 原先三份 TXT 没有来源字段：从当前工作簿反查，并排除同 URL 的其他 ID。
    for row in rows:
        if row["headline_id"] in ids and normalize_url(row["url"]):
            urls.add(normalize_url(row["url"]))
    return ids, urls, bodies


def grams(text):
    words = re.findall(r"\w+", text.lower())
    return {tuple(words[i:i + 5]) for i in range(len(words) - 4)}


def similar_to(body, bodies):
    a = grams(body)
    for ident, old in bodies.items():
        b = grams(old)
        intersection = len(a & b)
        union = len(a | b)
        minimum = min(len(a), len(b))
        if union and intersection / union > 0.65:
            return ident
        if minimum and intersection >= 50 and intersection / minimum > 0.8:
            return ident
    return None


def title_words(text):
    ignored = {"the", "a", "an", "and", "of", "to", "in", "on", "at", "for", "s", "read", "next", "watch", "live"}
    return set(re.findall(r"[\w]+", text.lower())) - ignored


def validate_article(row, article):
    expected, actual = title_words(row["headline"]), title_words(article.get("title", ""))
    if not actual:
        raise ValueError("missing_title")
    if len(expected) >= 3 and len(actual) < 3:
        raise ValueError("generic_title_needs_review")
    if len(expected) >= 3 and len(expected & actual) / min(len(expected), len(actual)) < 0.45:
        raise ValueError("title_mismatch_needs_review")
    expected_date = row.get("earliest_published_et")
    page_date = article.get("published")
    if expected_date and page_date:
        try:
            left = datetime.fromisoformat(str(expected_date)[:10]).date()
            right = datetime.fromisoformat(str(page_date)[:10]).date()
        except ValueError:
            raise ValueError("unreadable_date_needs_review")
        if abs((left - right).days) > 7:
            raise ValueError("date_mismatch_needs_review")
    elif expected_date:
        raise ValueError("missing_publication_date_needs_review")
    if len(" ".join(article.get("body", [])).split()) < 50:
        raise ValueError("too_short_needs_review")


class SafeRedirect(urllib.request.HTTPRedirectHandler):
    def __init__(self, allowed_hosts):
        self.allowed_hosts = allowed_hosts

    def redirect_request(self, req, fp, code, msg, headers, newurl):
        normalized = normalize_url(newurl)
        if not normalized or (self.allowed_hosts is not None and urllib.parse.urlsplit(normalized).hostname not in self.allowed_hosts):
            raise ValueError("redirect_to_unreviewed_source")
        if req.full_url.startswith("https:") and normalized.startswith("http:"):
            raise ValueError("insecure_redirect")
        return super().redirect_request(req, fp, code, msg, headers, normalized)


def fetch_html(url, allowed_hosts, timeout=30, *, system_proxy=False):
    if allowed_hosts is not None and urllib.parse.urlsplit(url).hostname not in allowed_hosts:
        raise ValueError("unreviewed_source")
    opener = public_opener(SafeRedirect(allowed_hosts), system_proxy=system_proxy)
    request = urllib.request.Request(url, headers={
        "User-Agent": "NewsTextCollector/1.0 (personal non-commercial research)",
        "Accept": "text/html,application/xhtml+xml",
    })
    with opener.open(request, timeout=timeout) as response:
        if response.headers.get_content_type() not in {"text/html", "application/xhtml+xml"}:
            raise ValueError("not_html")
        content = response.read(10 * 1024 * 1024 + 1)
        if len(content) > 10 * 1024 * 1024:
            raise ValueError("page_too_large")
        return content, response.geturl()


def make_text(row, article, source):
    metadata = [article["title"], "", "Headline ID: " + row["headline_id"],
                "Source URL: " + row["url"]]
    if article.get("author"):
        metadata.append("Author: " + str(article["author"]))
    if article.get("published"):
        metadata.append("Published: " + str(article["published"]))
    footer = "Personal, non-commercial study copy."
    if source.get("terms_url"):
        footer += " Publisher terms: " + source["terms_url"]
    if source.get("mode") == "generic":
        footer += " Generic extraction; completeness requires review."
    return ("\n".join(metadata) + "\n\n" + "\n\n".join(article["body"]) + "\n\n"
            + "\n".join(article.get("copyright", [])) + "\n\n" + footer + "\n")


def classify(row, ids, urls, sources, start, end):
    if row["excel_row"] < start or (end and row["excel_row"] > end):
        return "outside_row_range"
    ident = row["headline_id"]
    if not ID_RE.fullmatch(ident):
        return "invalid_headline_id"
    if ident in ids:
        return "existing_id"
    url = normalize_url(row["url"])
    if not url:
        return "invalid_url"
    if url in urls:
        return "duplicate_url"
    if row.get("match_quality") and row["match_quality"] != "high":
        return "source_match_needs_review"
    if sources is not None and urllib.parse.urlsplit(url).hostname not in sources:
        return "unreviewed_source"
    return "ready"


def run(args):
    if getattr(args, "probe_all", False):
        from check_urls import probe_all
        return probe_all(args)
    project = Path(args.project_dir).expanduser().resolve()
    workbook = Path(args.workbook).expanduser().resolve() if args.workbook else project / "news_urls(1).xlsx"
    output = Path(args.output).expanduser().resolve() if args.output else project / "New Output"
    rows = read_rows(workbook, args.sheet)
    if args.all_sources:
        sources = None
        print("通用模式：不读取 sources.json，尝试任何有效公开网址；正文质量仍需校对。")
    else:
        config = json.loads((HERE / "sources.json").read_text(encoding="utf-8"))
        sources = {host: item for item in config["sources"] for host in item["hosts"]}
        print("配置模式：仅尝试 sources.json 内的网站；加 --all-sources 可启用通用模式。")
    ids, urls, bodies = existing_state(project, output, rows)
    selected = [r for r in rows if r["excel_row"] >= args.start_row and (not args.end_row or r["excel_row"] <= args.end_row)]
    print("已有新闻：{} 篇；本次按 Excel 第 {} 行起依次检查。".format(len(ids), args.start_row))
    # 预览不联网，不创建新闻 TXT；完整行序/跳过原因写 CSV。
    stamp = datetime.now().astimezone().strftime("%Y%m%d_%H%M%S_%f")
    logdir = HERE / "runs"
    logdir.mkdir(exist_ok=True)
    if not args.download:
        report = logdir / ("preview_" + stamp + ".csv")
        preview, counts = [], {}
        for row in selected:
            status = classify(row, ids, urls, sources, args.start_row, args.end_row)
            counts[status] = counts.get(status, 0) + 1
            preview.append({k: row[k] for k in ("excel_row", "headline_id", "headline", "url")})
            preview[-1]["status"] = status
        with report.open("x", encoding="utf-8-sig", newline="") as f:
            writer = csv.DictWriter(f, fieldnames=["excel_row", "headline_id", "headline", "url", "status"])
            writer.writeheader(); writer.writerows(preview)
        print("预览统计：" + json.dumps(counts, ensure_ascii=False))
        if counts.get("ready", 0) == 0:
            print("当前规则下可自动尝试的新条目为0；请先核对其他来源并补充对应提取规则。")
        print("前 {} 个可尝试的 ID：".format(args.limit))
        for row in [r for r in preview if r["status"] == "ready"][:args.limit]:
            print("  第 {} 行  {}".format(row["excel_row"], row["headline_id"]))
        print("预览清单：" + str(report))
        print("未联网。确认后加 --download 才会下载。")
        return 0
    output.mkdir(parents=True, exist_ok=True)
    journal = logdir / ("download_" + stamp + ".jsonl")
    manifest_dir = output / "provenance" / "cursor_runs"
    manifest_dir.mkdir(parents=True, exist_ok=True)
    saved, attempted, blocked, seen = [], 0, set(), set()
    outcomes = Counter()
    last_request = 0.0
    stop_reason = "end_of_rows"
    with journal.open("x", encoding="utf-8") as log:
        try:
            for row in selected:
                status = classify(row, ids, urls, sources, args.start_row, args.end_row)
                url = normalize_url(row["url"])
                host = urllib.parse.urlsplit(url).hostname if url else ""
                record = {k: row[k] for k in ("excel_row", "headline_id", "headline", "url")}
                if status == "ready" and url in seen:
                    status = "url_already_attempted_this_run"
                if status == "ready" and host in blocked:
                    status = "host_blocked_this_run"
                if status == "ready":
                    if attempted >= args.max_requests:
                        stop_reason = "request_limit"; break
                    wait = args.delay - (time.monotonic() - last_request)
                    if wait > 0:
                        time.sleep(wait)
                    attempted += 1; seen.add(url)
                    try:
                        fetch_options = {"system_proxy": True} if args.system_proxy else {}
                        raw, final_url = fetch_html(url, None if sources is None else set(sources), args.timeout, **fetch_options)
                        try:
                            article = extract_article(raw, final_url)
                        except ExtractionError as exc:
                            if not args.all_sources or exc.code not in {
                                "unsupported_domain", "article_body_not_found", "empty_article", "incomplete_article"
                            }:
                                raise
                            article = extract_generic_article(raw, final_url)
                        source = {"mode": "generic", "terms_url": ""} if sources is None else sources[host]
                        validate_article(row, article)
                        body = "\n\n".join(article["body"])
                        duplicate = similar_to(body, bodies)
                        if duplicate:
                            status = "duplicate_text"; record["duplicate_of"] = duplicate
                        else:
                            path = output / (row["headline_id"] + ".txt")
                            text = make_text(row, article, source)
                            # 独占创建，避免覆盖已有文件；中断留下的半文件会清理。
                            with path.open("x", encoding="utf-8") as f:
                                try:
                                    f.write(text); f.flush()
                                except BaseException:
                                    f.close(); path.unlink(missing_ok=True); raise
                            status = "saved"
                            ids.add(row["headline_id"]); urls.add(url); urls.add(normalize_url(final_url))
                            bodies[row["headline_id"]] = body
                            record.update(title=article["title"], author=article.get("author"),
                                          publication_date=article.get("published"), final_url=final_url,
                                          path=str(path), body_words=len(body.split()), method=article["method"],
                                          warnings=article.get("warnings", []),
                                          copyright=article.get("copyright", []), terms_url=source.get("terms_url", ""),
                                          source_mode="generic" if sources is None else "configured",
                                          sha256=hashlib.sha256(text.encode()).hexdigest())
                            saved.append(record)
                    except urllib.error.HTTPError as e:
                        status = "http_" + str(e.code)
                        if e.code in (401, 403, 429):
                            blocked.add(host)
                        if e.fp is not None:
                            e.close()
                    except (ExtractionError, ValueError, OSError, urllib.error.URLError) as e:
                        status = "failed"
                        record["reason"] = str(e)[:500]
                        if isinstance(e, ExtractionError) and e.code == "challenge":
                            blocked.add(host)
                    finally:
                        last_request = time.monotonic()
                record["status"] = status
                outcomes[status] += 1
                log.write(json.dumps(record, ensure_ascii=False, default=str) + "\n"); log.flush()
                if status not in {"existing_id", "duplicate_url", "unreviewed_source", "source_match_needs_review", "invalid_url", "invalid_headline_id"}:
                    print("第 {} 行 {}: {}".format(row["excel_row"], row["headline_id"], status), flush=True)
                if len(saved) >= args.limit:
                    stop_reason = "target_reached"; break
        except KeyboardInterrupt:
            stop_reason = "interrupted"
            print("已停止，已成功保存的文件保留。")
    manifest = manifest_dir / ("run_" + stamp + ".json")
    manifest.write_text(json.dumps({"created_at": datetime.now().astimezone().isoformat(),
                                   "workbook": str(workbook), "sheet": args.sheet, "requested": args.limit,
                                   "saved_count": len(saved), "request_count": attempted, "stop_reason": stop_reason,
                                   "journal": str(journal), "articles": saved}, ensure_ascii=False, indent=2), encoding="utf-8")
    print("实际新增 {} / 目标 {}；请求 {} 次。结束原因：{}".format(len(saved), args.limit, attempted, stop_reason))
    print("处理统计：" + json.dumps(outcomes, ensure_ascii=False))
    if attempted == 0 and outcomes["unreviewed_source"]:
        print("未联网：{} 行因网站未列入 sources.json 被跳过。请在原命令中加 --all-sources，启用通用抓取。".format(outcomes["unreviewed_source"]))
    print("逐条记录：" + str(journal)); print("来源清单：" + str(manifest))
    return 0 if len(saved) == args.limit else 2


def parser():
    p = argparse.ArgumentParser(description=__doc__)
    p.add_argument("--project-dir", default=str(HERE.parent), help="包含原Excel、Current Output和New Output的文件夹")
    p.add_argument("--workbook", help="可选：Excel绝对路径")
    p.add_argument("--sheet", default="headline_urls")
    p.add_argument("--output", help="默认项目目录下New Output")
    group = p.add_mutually_exclusive_group()
    group.add_argument("--download", action="store_true", help="实际联网下载；默认仅预览")
    group.add_argument("--dry-run", action="store_true", help="只检查工作簿和已有文件")
    group.add_argument("--probe-all", action="store_true", help="绕过sources.json检查所有有效公开URL，只记录状态/标题，不保存正文")
    p.add_argument("--all-sources", action="store_true", help="预览/下载时不读取sources.json；未配置网站使用通用正文提取")
    p.add_argument("--system-proxy", action=argparse.BooleanOptionalAction, default=True,
                   help="默认沿用本机HTTP(S)代理；--no-system-proxy使用直连")
    p.add_argument("--workers", type=int, default=4, help="probe-all并发网站数，1–8；同站串行限速")
    p.add_argument("--limit", type=int, default=20, help="目标新增成功数，不是Excel行数")
    p.add_argument("--max-requests", type=int, default=40, help="最多尝试的不同初始URL数量，不包含重定向跳数")
    p.add_argument("--start-row", type=int, default=2, help="Excel实际行号，包含表头")
    p.add_argument("--end-row", type=int)
    p.add_argument("--delay", type=float, default=3.0)
    p.add_argument("--timeout", type=float, default=30.0, help="每次socket操作超时秒数，不是整个任务的时长上限")
    return p


if __name__ == "__main__":
    args = parser().parse_args()
    if args.limit < 1 or args.max_requests < 1 or args.start_row < 2 or args.delay < 1 or args.timeout <= 0 or not 1 <= args.workers <= 8:
        raise SystemExit("limit/max-requests需为正数；start-row至少2；delay至少1秒；timeout为正数。")
    if args.end_row and args.end_row < args.start_row:
        raise SystemExit("end-row不能小于start-row。")
    try:
        raise SystemExit(run(args))
    except (ValueError, OSError) as exc:
        raise SystemExit("无法运行：" + str(exc))
